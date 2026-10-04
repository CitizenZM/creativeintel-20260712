import json, os, re
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from plan import PRODUCTS, ADS, TV_DISCLAIMER

OUT = os.path.expanduser("~/Downloads/TCL-黑五网一-15秒广告策划-2026-10-04.docx")
assets = json.load(open("assets.json"))
doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.6)
st = doc.styles["Normal"]; st.font.name = "PingFang SC"; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
for n in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
    h = doc.styles[n]; h.font.name = "PingFang SC"; h.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC"); h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

def p(t, b=False, size=None, color=None):
    para = doc.add_paragraph(); r = para.add_run(t); r.bold = b
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    return para
def bl(t, pre=None):
    para = doc.add_paragraph(style="List Bullet")
    if pre: para.add_run(pre).bold = True
    para.add_run(t)
def table(head, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(head)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                for r in pp.runs: r.font.size = Pt(8.5)
    doc.add_paragraph()

doc.add_heading("TCL 黑五 / 网一 15 秒视频广告策划", 0)
p("QM7L · QM8L · NXTPAPER 14 — 每个产品 3 个创意，共 9 条 | 核心卖点：免费专业安装 | 2026-10-04 | CreativeIntel", color=(0x55, 0x5F, 0x70))

doc.add_heading("一、结论摘要", 1)
bl("官网 Best Seller 页（us.tcl.com/pages/best-seller）与电视按销量排序一致：QM7L 第 1、QM8L 第 2；NXTPAPER 14 是平板系列首位。", "选品：")
bl("TCL 官网已上线 FREE Professional TV Installation：购买合格电视可免费加购，含送货就位、专业安装、标准挂墙；挂架硬件与非标准工作另计；由 TCL 授权第三方执行，视地区而定。QM7L、QM8L 均在合格清单内。", "安装卖点已核实：")
bl("NXTPAPER 14 是平板，官网安装服务只针对电视。策划把“包安装”转译为“开箱即用 / 不用再装显示器”，请确认是否接受，或改为“买电视送安装 + 平板加购”的组合（需 TCL 提供该组合活动）。", "需你决定：")
bl("黑五（11/27）、网一（11/30）价格尚未公布，全部标为 [占位]；当前官网价已有较大折扣（QM7L 55\" 立省 $400，98\" 立省 $1,500；QM8L 65\" 立省 $1,000）。", "价格：")
bl("QM8L 主图带 NFL 商标、多张屏幕图含电影《Wicked》海报和 App 界面，均不可直接用于广告；屏幕内容一律本地合成自有画面。", "素材风险：")
bl("每条约 8 个镜头（单镜 ≤ 2 秒）+ 5 秒片尾卡；旁白 24–28 词（约 10 秒）；动作镜头走 Kling 3.0，其余走 Veo 3.1 Lite；预计每条 $1.6–2.0，9 条约 $15–18。", "制作：")

doc.add_heading("二、产品研究", 1)
for pr in PRODUCTS:
    doc.add_heading(f"{pr['name']}（{pr['rank']}）", 2)
    p(f"价格：{pr['prices']}"); p(f"来源：{pr['url']}", color=(0x55, 0x5F, 0x70))
    table(["核心规格（官网）", "广告卖点（如何用镜头证明）"], [(s, pr["sell"][i] if i < len(pr["sell"]) else "") for i, s in enumerate(pr["specs"])] + [("", x) for x in pr["sell"][len(pr["specs"]):]], [8.5, 8.5])
doc.add_heading("免费专业安装：官网原文要点与广告免责声明", 2)
for t in ["合格机型：QM7L、QM8L、X11L、RM9L、X11K、115\" QM7K", "流程：下单时勾选 → 安装伙伴主动来电 → 预约上门；无需单独预约",
          "包含：送货与就位、专业安装、标准挂墙", "不含：挂架/额外硬件、非标准安装、入墙走线等额外工作（可能收费）", "执行：TCL 授权第三方；视地区而定"]: bl(t)
p("广告统一免责声明（片尾小字）：", b=True); p(TV_DISCLAIMER)
p("注意：VO 与字幕不得说“wall mount included”（挂架硬件不含），统一用“we install it free / professional installation on us”。", color=(0xB0, 0x3A, 0x2E))

doc.add_heading("三、官方素材清单（已下载 39 张原图）", 1)
p("存放：creativeintel-20260712/out/research/tcl-bfcm/<产品>/；每张已用视觉模型识别内容与可用性。")
def risk(a):
    s = (a.get("shows") or "") + " " + (a.get("ad_use") or "")
    if re.search(r"NFL", s): return "不可用：NFL 商标"
    if re.search(r"Wicked|YouTube|app", s, re.I): return "不可用：影视/App 界面版权"
    if a.get("has_text"): return "慎用：含文字，仅作片尾参考"
    return "可用：干净产品图（屏幕内容另行合成）"
table(["文件", "尺寸", "内容", "可用性"], [(a["file"], a["px"], a.get("shows", ""), risk(a)) for a in assets], [4, 2, 7, 4.5])
p("NXTPAPER 14 官网只有产品图，没有使用场景图；乐谱、学习、办公等场景需要 AI 生成（演员锁定 + 官方产品图作参考）。")

doc.add_heading("四、9 个 15 秒广告创意", 1)
p("统一结构：0–1.5 秒钩子 → 5 秒内出现“安装/开箱”回报 → 5–10 秒卖点证明（每个镜头一个运镜，单镜 ≤ 2 秒）→ 10–15 秒片尾卡（官方产品图本地合成 + 价格 + CTA + 免责声明）。开头和结尾用强转场；字幕为 1–4 个词的卖点关键词。")
for a in ADS:
    doc.add_heading(f"{a['id']}｜{a['title']}", 2)
    table(["节点", "受众", "钩子类型", "洞察"], [(a["event"], a["audience"], a["hook"], a["insight"])], [3, 4, 2.5, 8])
    table(["时间(s)", "画面", "运镜", "屏幕关键词", "渲染"], [list(s) for s in a["shots"]], [1.6, 8, 2.6, 3.3, 1.8])
    p("旁白（VO, en-US）：", b=True); p(f"“{a['vo']}”  — {len(a['vo'].split())} 词 ≈ {len(a['vo'].split())/2.6:.0f} 秒")
    p("为什么能打：", b=True); p(a["why"])
    if PRODUCTS[[x["id"] for x in PRODUCTS].index(a["product"])]["install"]:
        p("片尾免责声明：使用第二节统一声明。", color=(0x55, 0x5F, 0x70))

doc.add_heading("五、制作方案（确认后执行）", 1)
for t in ["导演 v2 自动执行：卖点→镜头证明、每镜运镜、≤2 秒构图、开头/结尾强转场、胶片调色（去塑料感）。",
          "人物大动作镜头（安装师傅搬抬、人物跳起、甩身坐下）→ Kling 3.0 Std（$0.42/条，按 5 秒计费）；产品与氛围镜头 → Veo 3.1 Lite（$0.12/4 秒）。",
          "演员锁定：每条广告先生成一张演员参考图，所有人物镜头保持同一张脸；电视/平板以官方干净产品图为参考。",
          "屏幕内容：一律本地合成自有授权画面（风景、通用赛车、通用橄榄球无标识），不用 NFL / 影视 / App 界面。",
          "交付：9:16 主版 + 4:5 / 1:1 / 16:9，及 10 秒剪版；每条附 QC 13 项报告与导演评分。",
          "预算：每条约 $1.6–2.0（OpenRouter），9 条约 $15–18；单条超过 120¢ 上限，需你批准后编译。"]: bl(t)

doc.add_heading("六、需要你确认", 1)
for t in ["9 个创意方向是否通过；每个产品先做哪一条（建议：QM7L-A、QM8L-B、NXT-A）。",
          "NXTPAPER 14 的“包安装”转译为“开箱即用 / 不用装显示器”是否可以。",
          "NXTPAPER 14 套装是否包含 T-Pen 与保护套（官网文案前后不一致）。",
          "黑五/网一的最终价格与活动起止日期（替换 [占位]）。",
          "是否允许出现通用橄榄球比赛画面（无任何 NFL 标识）。",
          "演员形象偏好（年龄/族裔）与是否需要真人口播。"]: bl(t)

doc.add_heading("资料来源", 1)
for t in ["https://us.tcl.com/pages/best-seller", "https://us.tcl.com/products/qm7l-series-sqd-mini-led-4k-uhd-hdr-smart-tv-with-google-tv",
          "https://us.tcl.com/products/qm8l-series-sqd-mini-led-4k-uhd-hdr-smart-tv-with-google-tv", "https://us.tcl.com/products/nxtpaper-14",
          "https://www.tomsguide.com/tvs/4k-tvs/tcl-qm8l-vs-qm7l-which-sqd-mini-led-tv-tested-better-in-our-lab",
          "https://www.geeky-gadgets.com/tcl-nxtpaper-14/", "https://www.itechguides.com/tcl-nxtpaper-14-review-a-huge-matte-screen-that-makes-reading-easier/",
          "https://www.techtimes.com/articles/315664/20260404/tcl-nxtpaper-14-go-tablet-musicians-digital-sheet-music-users.htm"]: bl(t)
doc.save(OUT); print(OUT)
